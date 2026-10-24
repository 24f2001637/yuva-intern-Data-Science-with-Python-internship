"""Build Weeks 4 and 5 using the approved Week 3 report as a design template.

Run analysis first. Requires python-docx; does not regenerate statistical results.
"""
from pathlib import Path
import json
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT/'weeks/week03/Week3_Statistical_Analysis_Hypothesis_Testing.docx'
M = json.loads((ROOT/'weeks/week04/outputs/results.json').read_text())
SELECTED = next(x for x in M['metrics'] if x['model']==M['selected_model'])
NAMES = {4:'Week4_Machine_Learning_Development_Evaluation.docx',
         5:'Week5_Comprehensive_Project_Report.docx'}

def start(week, title1, title2, analysis):
    d=Document(TEMPLATE)
    body=d.element.body
    # Keep the approved cover, cover section break, final section settings,
    # image relationships, and section-specific full-page background.
    for e in list(body)[2:-1]:body.remove(e)
    c=d.tables[0].cell(0,0)
    changes={1:f'W  E  E  K  -  {week}     R  E  P  O  R  T',2:title1,3:title2,
        4:f'Week {week} Deliverable — Titanic Survival Analysis',
        8:'Tools: Python 3.12 · pandas · scikit-learn · Matplotlib',9:analysis,
        12:f'YuvaIntern Data Science Apprenticeship                                        Week {week} of 5'}
    for i,text in changes.items():
        p=c.paragraphs[i]; rp=deepcopy(p.runs[0]._r.rPr)
        for r in list(p.runs):p._p.remove(r._r)
        r=p.add_run(text);r._r.insert(0,rp)
        if i in (2,3):r.font.size=Pt(32)
    h=d.sections[1].header.paragraphs[0]
    for r in list(h.runs):h._p.remove(r._r)
    r=h.add_run(title1+' '+title2);r.font.name='Calibri';r.font.size=Pt(9)
    r.font.color.rgb=RGBColor.from_string('6878A0');r.italic=True
    d.styles['Normal'].font.name='CMU Serif';d.styles['Normal'].font.size=Pt(10.5)
    d.styles['Normal'].paragraph_format.line_spacing=1.12
    d.styles['Normal'].paragraph_format.space_after=Pt(6)
    return d

def chapter(d,title,first=False):
    p=d.add_paragraph(title,'Heading 1')
    for r in p.runs:r.font.name='CMU Serif';r.font.color.rgb=RGBColor.from_string('0A1628')
    p.paragraph_format.page_break_before=not first
    p.paragraph_format.space_before=Pt(0)
    return p

def p(d,text):
    q=d.add_paragraph(text)
    q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    q.paragraph_format.first_line_indent=Inches(.25)
    return q

def sub(d,text):
    q=d.add_paragraph(text,'Heading 2')
    for r in q.runs:r.font.name='CMU Serif';r.font.color.rgb=RGBColor.from_string('0A1628')
    return q

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=widths or [6.1/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,s in enumerate(headers):t.rows[0].cells[i].text=str(s)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trpr.append(no)
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge)
                for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
                borders.append(e)
            pr.append(borders)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE7EF' if ri==0 else ('F4F8FC' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            mar=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(2);q.paragraph_format.line_spacing=1.05
                for r in q.runs:r.font.name='CMU Serif';r.font.size=Pt(9.5);r.bold=ri==0
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def figure(d,path,caption,width=5.7):
    q=d.add_paragraph();q.paragraph_format.keep_with_next=True
    q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.add_run().add_picture(str(ROOT/path),width=Inches(width))
    d.add_paragraph(caption,'Caption')

def code(d,text):
    q=d.add_paragraph();q.paragraph_format.space_after=Pt(6)
    q.paragraph_format.line_spacing=1.05
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F4F8FC');q._p.get_or_add_pPr().append(sh)
    r=q.add_run(text.strip());r.font.name='DejaVu Sans Mono';r.font.size=Pt(8)
    r.font.color.rgb=RGBColor.from_string('1A2B40')

def references(d,items):
    for i,text in enumerate(items,1):
        q=d.add_paragraph(f'[{i}] {text}');q.paragraph_format.line_spacing=1.05
        for r in q.runs:r.font.size=Pt(9)

def week4():
    d=start(4,'Machine Learning','Development and Evaluation','Analysis: 3 model comparisons and 4 evaluation figures')
    chapter(d,'1.  Introduction and modelling objective',True)
    p(d,'This report develops and evaluates a basic classification model for Titanic passenger survival. I use the same 891-record source snapshot as the previous weeks, preserving continuity from exploratory analysis to statistical testing and prediction. The target is survived, where 1 means survival and 0 means non-survival. The objective is an educational comparison of predictive methods on historical passenger records, rather than an operational evacuation system.')
    p(d,f'Logistic regression was selected using five-fold cross-validation within a 712-record development set. On the remaining 179 records it achieved {SELECTED["accuracy"]:.1%} accuracy, {SELECTED["roc_auc"]:.3f} ROC AUC and {SELECTED["f1"]:.3f} F1. It outperformed the majority-class baseline and the constrained decision tree in this experiment. Its errors remain consequential: 21 of 69 actual survivors were classified as non-survivors, and subgroup performance was uneven.')
    sub(d,'Continuity with the earlier analysis')
    p(d,'Weeks 1 and 2 supplied descriptive patterns; Week 3 tested associations with sex, class, age, childhood status and family group. Those findings motivate candidate predictors but do not establish causation. In particular, family size has a non-linear pattern, and the age comparison is sensitive to missing-value treatment. Predictive usefulness is assessed here through cross-validation rather than inferred from statistical significance alone.')
    p(d,'The full dataset was explored in earlier weeks before this split was created. Consequently, the test records were withheld from Week 4 model fitting and selection, but they are not a pristine external validation set. The scores are internal estimates for this particular historical dataset and may be optimistic after project-wide exploration.')
    table(d,['Design item','Specification'],[
        ['Prediction target','survived: 0 or 1'],['Source records','891; 342 survivors and 549 non-survivors'],
        ['Development / test','712 / 179 records; stratified split, random seed 42'],
        ['Model selection','Highest mean five-fold development ROC AUC'],
        ['Classification threshold','0.50, fixed before test evaluation'],
        ['Primary limitation','No independent external validation; related passengers may cross splits']],[1.65,4.45])

    chapter(d,'2.  Data preparation and leakage prevention')
    p(d,'I load the original CSV through the shared hash-validated loader. I retain all 891 records: repeated feature profiles do not identify duplicate people because this dataset contains no passenger names, ticket identifiers or family IDs. The globally imputed Week 1 EDA table is not used for model fitting. Instead, preprocessing is learned separately within each training fold and then applied to its validation fold.')
    table(d,['Input','Preparation and reason'],[
        ['Age','Median imputation within each training fold; a missingness indicator retains information about absent ages.'],
        ['Fare','log(1 + fare) reduces the long right tail and permits zero fares; scaled inside the pipeline.'],
        ['Passenger class','One-hot encoded so the model does not require equal effects between classes.'],
        ['Sex and embarked','Categorical encoding; mode imputation fitted only on training records.'],
        ['Family group','sibsp + parch + 1 grouped as solo, 2–4, or 5+; one-hot encoding represents a non-linear pattern.']],[1.25,4.85])
    p(d,'There are 177 missing ages and two missing embarkation values in the source data. Numerical inputs are standardized after imputation. Categorical encoding ignores previously unseen categories at prediction time, avoiding a runtime failure while highlighting the need to monitor new categories in any future application. Deterministic transformations such as log fare and family-group construction do not estimate parameters from the dataset.')
    sub(d,'Excluded variables')
    p(d,'The alive column is excluded because it directly restates the outcome and would cause target leakage. The class label duplicates pclass. The who, adult_male and alone fields are derived from retained characteristics and are omitted for a simpler, explicit feature set; redundancy itself is not automatically target leakage. Deck is excluded because 688 of 891 values are missing. sibsp and parch are used only to construct family group.')
    p(d,'A scikit-learn Pipeline contains imputation, encoding, scaling and the estimator. Cross-validation receives the entire pipeline, so validation records cannot influence fitted medians, modes or scaling statistics. The test set is not used for hyperparameter selection or threshold changes. This follows the pipeline and leakage-prevention practices in references [2] and [3].')

    chapter(d,'3.  Model selection and training')
    p(d,'I compare a majority-class baseline with two simple supervised classifiers. All models use the same split, folds and permitted input information. No hyperparameter search was performed: the settings below were fixed for an interpretable introductory comparison. The baseline predicts the development-set survival prevalence as a constant probability; at threshold 0.50 it predicts non-survival for everyone.')
    table(d,['Candidate','Settings and rationale'],[
        ['Majority baseline','DummyClassifier(strategy="prior"); establishes whether learned features add value.'],
        ['Logistic regression','C=1, default L2 regularization, max_iter=2000. A stable additive model that outputs class probabilities.'],
        ['Decision tree','max_depth=4, min_samples_leaf=10, random_state=42. Captures interactions and non-linear splits while limiting complexity.']],[1.45,4.65])
    p(d,'The development set contains 273 survivors and 439 non-survivors. Five shuffled stratified folds preserve approximate class balance. ROC AUC is the selection criterion because it measures ranking across thresholds. Accuracy and F1 are also recorded. Cross-validation standard deviations summarize fold variation; they are not confidence intervals and do not prove a statistically significant difference between models.')
    table(d,['Model','Train AUC¹','CV AUC','CV SD'],[[r['model'],f'{r["cv_train_auc"]:.3f}',f'{r["cv_auc_mean"]:.3f}',f'{r["cv_auc_sd"]:.3f}'] for r in M['metrics']],[2.5,1.2,1.2,1.2])
    d.add_paragraph('¹ Mean training-fold AUC from the same cross-validation procedure.','Caption')
    p(d,'Logistic regression has the highest mean development AUC (0.866 versus 0.853 for the tree) and is selected before test probabilities are examined. The difference is modest relative to fold variation. Each candidate is refitted on all 712 development records for the final comparison, but no candidate is refitted on the test set.')
    p(d,'The mean train-to-validation AUC gap is approximately 0.007 for logistic regression and 0.036 for the tree. The larger tree gap suggests more overfitting in this comparison. Logistic regression can still underfit interactions: its additive form does not explicitly model the sex-by-class interaction identified in Week 3. An interaction model is therefore a future candidate, not an improvement established by these results.')

    chapter(d,'4.  Overall model performance')
    figure(d,Path('weeks/week04/outputs/01_model_comparison.png'),'Figure 4.1  Development cross-validation and test ROC AUC. Error bars represent the standard deviation across five development folds, not test-set uncertainty.')
    table(d,['Model','Accuracy','Precision','Recall','F1'],[[r['model'],f'{r["accuracy"]:.3f}',f'{r["precision"]:.3f}',f'{r["recall"]:.3f}',f'{r["f1"]:.3f}'] for r in M['metrics']],[2.1,1,1,1,1])
    p(d,'The baseline achieves 61.5% accuracy while identifying no survivors. This illustrates why accuracy alone is insufficient on an imbalanced target. Precision is the fraction of predicted survivors who actually survived; recall is the fraction of actual survivors detected. F1 combines precision and recall. Baseline precision is undefined with no positive predictions and is reported as zero using an explicit zero_division setting.')
    p(d,'Logistic regression improves accuracy by 20.7 percentage points over the baseline and reaches 79.8% balanced accuracy. Its average precision is 0.795 and Brier score is 0.141, compared with 0.385 and 0.237 for the constant-probability baseline. Lower Brier scores indicate smaller squared probability errors; this comparison alone does not establish well-calibrated probabilities.')

    chapter(d,'5.  Confusion matrices and ranking quality')
    figure(d,Path('weeks/week04/outputs/02_confusion_matrices.png'),'Figure 4.2  Test confusion matrices at the fixed 0.50 threshold. Rows show observed outcomes; columns show predicted outcomes.')
    p(d,'For logistic regression, 99 passengers are correctly classified as non-survivors and 48 as survivors. There are 11 false positives and 21 false negatives. Thus 147 of 179 predictions are correct and 32 are incorrect, giving a 17.9% error rate. The tree identifies fewer survivors at the same threshold, with recall of 55.1% versus 69.6% for logistic regression.')
    figure(d,Path('weeks/week04/outputs/03_roc_pr.png'),'Figure 4.3  Test ROC and precision–recall curves for all candidates. ROC AUC assesses ranking; precision–recall focuses on survivor detection. The baseline average precision equals the test survival prevalence.')
    p(d,'The selected model has test ROC AUC 0.845. These curves describe threshold trade-offs without choosing a new threshold from the test set. Any future threshold selection should use development validation predictions and a stated error-cost objective, followed by evaluation on new independent data.')

    chapter(d,'6.  Error analysis and uncertainty')
    figure(d,Path('weeks/week04/outputs/04_subgroups.png'),'Figure 4.4  Survivor recall by sex and passenger class for the selected model. Denominators are shown because small groups can produce unstable rates. Groups overlap across the two attributes.')
    p(d,'All 45 female survivors in the test split are identified, but only 3 of 24 male survivors are identified. Female false-positive rate is 9/16 (56.3%), whereas male false-positive rate is 2/94 (2.1%). Overall accuracy therefore conceals substantial differences in error type. These are descriptive results from one small split, not definitive population estimates or evidence of fair performance.')
    p(d,'By passenger class, survivor recall is 68.0% for first class (17/25), 85.0% for second class (17/20), and 58.3% for third class (14/24). The sample does not support reliable intersectional conclusions for every sex-by-class subgroup. Before broader use, I would report larger subgroup samples and uncertainty rather than interpret these percentages in isolation.')
    p(d,'A 2,000-resample percentile bootstrap of test passengers gives an approximate 95% accuracy interval of 76.5%–87.2% and ROC AUC interval of 0.775–0.906. These intervals are conditional on this fitted model and split. They do not account for model-selection variability, earlier full-data exploration, or dependence between relatives. They should not be presented as guaranteed external performance bounds.')

    chapter(d,'7.  Limitations and improvement plan')
    p(d,'This historical sample comes from a single event. It does not represent modern travel, evacuation systems or customer behaviour. Many potentially relevant variables are absent, including movement routes, access to assistance and complete family identifiers. Related passengers may appear in both training and test partitions, which can make ordinary row-wise validation optimistic. Group-aware splitting requires reliable group identifiers that are not available here.')
    p(d,'The preprocessing pipeline avoids fitting transformations on validation or test records, but it cannot reverse the fact that the full dataset informed earlier exploration. New independent records are needed for a stronger claim about generalization. The model learns patterns associated with observed survival; it does not identify the effect of changing class, fare or any other passenger attribute.')
    table(d,['Next improvement','How I would assess it'],[
        ['Explicit interactions','Compare a pre-specified sex-by-class model using development cross-validation.'],
        ['Regularization and tree tuning','Use a limited grid inside development CV; reserve independent evaluation data.'],
        ['Missing-data sensitivity','Compare fold-fitted imputation strategies and retain missingness indicators.'],
        ['Calibration and threshold','Use out-of-fold probabilities; choose metrics and error costs before evaluation.'],
        ['Group-aware external validation','Acquire family or ticket linkage; assess performance on independent groups.']],[1.8,4.3])
    p(d,'The current model is suitable for demonstrating a reproducible workflow, inspecting errors and generating research questions. It is not validated for assigning rescue priority or allocating services. Historical sex and class associations must not be converted directly into operational decisions.')
    sub(d,'Workload plan')
    p(d,'A proposed 32-hour allocation is: data audit and feature preparation, 6 hours; pipeline and candidate development, 8 hours; cross-validation and evaluation, 7 hours; error analysis, 5 hours; and documentation and review, 6 hours. This is a planning allocation within the assigned 30–35 hours, not a claim of recorded time spent.')

    chapter(d,'8.  Python implementation and reproducibility')
    p(d,'The complete executed script is weeks/week04/analysis.py. It resolves paths from its own location and loads the pinned CSV through src/titanic_data.py. The following excerpts show the key preparation and selection logic; plotting, bootstrap and output-writing code remain in the complete script.')
    code(d,'''numeric = Pipeline([
    ('impute', SimpleImputer(strategy='median',
                            add_indicator=True)),
    ('scale', StandardScaler())])
categorical = Pipeline([
    ('impute', SimpleImputer(strategy='most_frequent')),
    ('encode', OneHotEncoder(handle_unknown='ignore',
                            sparse_output=False))])
prep = ColumnTransformer([
    ('numeric', numeric, ['age', 'log_fare']),
    ('categorical', categorical,
     ['pclass', 'sex', 'embarked', 'family_group'])])
pipe = Pipeline([('prepare', prep), ('model', model)])''')
    code(d,'''train, test = train_test_split(raw.index, test_size=.2,
    stratify=y, random_state=42)
cv = StratifiedKFold(n_splits=5, shuffle=True,
                    random_state=42)
result = cross_validate(pipe, x.loc[train], y.loc[train],
    cv=cv, scoring={'auc': 'roc_auc',
                   'accuracy': 'accuracy', 'f1': 'f1'},
    return_train_score=True)
# Repeat for the three fixed candidates, then select by CV.
selected = max(models,
    key=lambda n: np.mean(scores[n]['test_auc']))
# Each pipeline is fitted on development records only.
p = fitted[selected].predict_proba(x.loc[test])[:, 1]
prediction = (p >= .5).astype(int)''')
    p(d,'Install requirements.txt in a Python 3.12 environment and run python weeks/week04/analysis.py. Results include results.json, model_metrics.csv, test_predictions.csv, subgroup_metrics.csv, split.csv and four PNG figures. The split manifest uses zero-based source row numbers, not unique passenger identifiers. Package versions and the dataset SHA256 are recorded in results.json.')

    chapter(d,'9.  Conclusions and references')
    p(d,'The experiment meets the Week 4 objective by preparing features, comparing a baseline and two basic algorithms, selecting a model on development data, and evaluating its predictions with multiple metrics and visualizations. Logistic regression is the selected candidate, with 82.1% test accuracy and 0.845 ROC AUC. Its advantage over the baseline is clear in this split, but the comparison does not establish universal superiority over trees or other classifiers.')
    p(d,'The most important practical lesson is that a good aggregate score does not remove the need to inspect errors. The selected model misses 21 actual survivors, all male in this split, and its female false-positive rate is high. Together with historical sampling, missing covariates and prior full-data exploration, these findings limit the conclusions to an educational internal benchmark. Week 5 synthesizes these results with the earlier statistical evidence and turns them into proportionate research and process recommendations.')
    references(d,[
        'Seaborn example-data repository, Titanic CSV. https://github.com/mwaskom/seaborn-data/blob/master/titanic.csv',
        'scikit-learn, Common pitfalls and recommended practices. https://scikit-learn.org/stable/common_pitfalls.html',
        'scikit-learn, Cross-validation: evaluating estimator performance. https://scikit-learn.org/stable/modules/cross_validation.html',
        'scikit-learn, Classification metrics. https://scikit-learn.org/stable/modules/model_evaluation.html',
        'Project evidence: supplied Week 1 and Week 2 reports; Week 3 report and outputs/results.json; Week 4 analysis.py and outputs.'])
    sub(d,'Execution environment')
    p(d,'The analysis was executed with Python 3.12.14, scikit-learn 1.8.0, pandas 2.2.3 and NumPy 2.3.5. Seeds are fixed at 42. Numerical values shown in the report are rounded from saved outputs. No elapsed internship hours, external deployment or operational benefits are claimed.')
    d.save(ROOT/'weeks/week04'/NAMES[4])

def week5():
    d=start(5,'Comprehensive Project Report','and Strategic Recommendations','Analysis: integrated findings, model evaluation and action plan')
    chapter(d,'1.  Executive summary',True)
    p(d,'This report brings together the Titanic data science project completed across Weeks 1–4 and presents evidence-based recommendations for reviewers and future project teams. The project progresses from data acquisition and cleaning to visual storytelling, statistical analysis and machine learning. Its central conclusion is that passenger survival is strongly associated with recorded sex, class and family group, while predictive accuracy alone provides an incomplete account of model quality.')
    p(d,'The source contains 891 passenger records and 342 survivors, a survival rate of 38.4%. All records are retained after reviewing the meaning of repeated profiles. Statistical tests support several associations, but the mean-age difference is small and sensitive to missing-data treatment. A class-only adjustment does not eliminate the embarkation-port association in the specified model, correcting an overstatement in the earlier visual narrative.')
    p(d,'Week 4 selected logistic regression by development cross-validation. It achieved 82.1% accuracy and 0.845 ROC AUC on a 179-record internal test split, compared with 61.5% accuracy for the majority baseline. It nevertheless missed 21 of 69 survivors. Survivor recall differed sharply by sex, illustrating the need to report error types and subgroup denominators alongside aggregate metrics.')
    table(d,['Priority','Recommended decision'],[
        ['Immediate','Preserve the source snapshot, explicit preprocessing rules, executable code and split manifest.'],
        ['Before broader use','Obtain independent validation data and audit subgroup errors and calibration.'],
        ['Future research','Collect family linkage and operational variables; evaluate proposed interventions separately.']],[1.45,4.65])
    p(d,'The expected value is improved analytical reliability, clearer communication and better-designed follow-up research. The dataset does not establish a modern safety intervention, customer retention strategy or financial return. Any contemporary operational proposal must be treated as a hypothesis requiring relevant new data and domain review.')
    p(d,'This is the final report of the five-week internship. It integrates the completed technical work and provides a documented handover of findings, code, limitations and proposed next steps.')

    chapter(d,'2.  Project scope and data foundation')
    p(d,'The project asks which passenger characteristics are associated with Titanic survival, whether those patterns withstand statistical examination, and how well simple models predict the recorded outcome. The intended audience includes the internship evaluator, technical reviewers and readers interested in responsible interpretation of observational data. It is a historical learning project, not a validated decision system.')
    table(d,['Stage','Contribution to the project'],[
        ['Week 1','Public data acquisition, missing-value audit, preprocessing and initial exploratory charts.'],
        ['Week 2','A visual narrative examining sex and class, age, family size, fares and embarkation.'],
        ['Week 3','Five hypothesis tests with Holm adjustment, effect sizes, confidence intervals and follow-up models.'],
        ['Week 4','Pipeline-based preprocessing, baseline and classifier comparison, internal test evaluation and error analysis.'],
        ['Week 5','Synthesis, limitations, recommendations, success measures and reproducibility handover.']],[1.1,5.0])
    p(d,'The dataset is the Seaborn Titanic CSV snapshot with 891 rows and 15 columns. Its SHA256 is retained in the repository to detect accidental changes. The source contains 177 missing ages, 688 missing deck values and two missing embarkation values. It has no names or ticket IDs, so repeated feature profiles cannot be treated as proof of duplicate passengers.')
    p(d,'The original Week 1 and Week 2 reports are preserved as historical submissions. Their source scripts were not supplied, so the repository contains explicitly labelled reconstructed analyses. Later reports document corrections rather than silently rewriting the original submissions. This distinction lets a reviewer trace both the project’s development and the currently supported conclusions.')
    sub(d,'Different preparation for different questions')
    p(d,'The EDA workflow may use imputed ages to display complete visual summaries, while preserving observed ages and missingness indicators. Week 3 age inference uses the 714 observed ages and treats imputation as a sensitivity check. Week 4 learns imputation and encoding only within training partitions. A single globally cleaned table is therefore not reused indiscriminately for all analytical purposes.')

    chapter(d,'3.  Descriptive findings and their interpretation')
    figure(d,Path('weeks/week03/outputs/rates.png'),'Figure 5.1  Survival proportions by recorded sex and passenger class from the audited Week 3 analysis. These are historical associations; the intervals depend on the working sampling assumptions.')
    p(d,'Female passengers have a recorded survival rate of 74.2% (233/314), compared with 18.9% (109/577) for male passengers. First-class survival is 63.0% (136/216), second-class survival is 47.3% (87/184), and third-class survival is 24.2% (119/491). These large descriptive differences justify further examination, but do not identify the mechanisms responsible for them.')
    figure(d,Path('weeks/week03/outputs/groups.png'),'Figure 5.2  Childhood status uses observed ages; family grouping uses all passengers. Small families of 2–4 have the highest survival proportion among the three family groups.')
    p(d,'Survival is 30.4% for solo travellers, 57.9% for families of 2–4, and 16.1% for families of 5 or more. The grouped relationship is an inverted U, correcting the earlier description of a U-shaped curve. These observations do not establish that family size itself caused the outcome or predict modern customer loyalty.')

    chapter(d,'4.  Statistical evidence and revised conclusions')
    table(d,['Question','Records','Holm p-value','Interpretation'],[
        ['Sex and survival','891','1.856 × 10⁻⁵⁸','Strong association'],
        ['Class and survival','891','1.820 × 10⁻²²','Strong association'],
        ['Mean age by survival','714','0.04119','Small, sensitive difference'],
        ['Under 13 and survival','714','0.004025','Association under assumptions'],
        ['Family group and survival','891','1.957 × 10⁻¹⁶','Non-linear grouped association']],[1.7,.65,1.2,2.55])
    p(d,'The five-test family uses a 0.05 significance threshold with Holm adjustment. These questions were informed by earlier exploration of the same data, so the results remain exploratory rather than preregistered confirmation. P-values express compatibility with specified null models under assumptions; they are not the probability that a hypothesis is true and do not measure practical importance.')
    p(d,'Observed-age survivors are on average 2.28 years younger than non-survivors. The Welch 95% interval for the difference is −4.47 to −0.09 years, and Hedges’ g is −0.157, a small standardized difference. The group-median-imputed sensitivity analysis gives an unadjusted p-value of approximately 0.0817. This change shows why missing-data decisions need explicit reporting.')
    sub(d,'Corrections that matter for decisions')
    p(d,'A likelihood-ratio comparison adding embarkation port to a class-only logistic model gives p≈0.0000778. Class alone therefore does not fully account for the observed port association in that specified model. This does not establish a causal port effect; omitted variables and model specification remain possible explanations. The earlier claim that the effect disappeared entirely was too strong.')
    p(d,'A sex-by-class interaction improves the specified Week 3 logistic model (likelihood-ratio statistic 28.79 on two degrees of freedom, p≈0.000000560). This supports investigating variation across combinations of attributes, rather than assuming one uniform association. The follow-up tests are exploratory and were not part of the five-test Holm family.')
    p(d,'Independence is a working assumption, not a verified property of these passenger records. Relatives may share outcomes and circumstances. Without reliable family IDs, ordinary confidence intervals and p-values cannot fully account for this clustering. No causal or population-wide operational conclusion follows from significance alone.')

    chapter(d,'5.  Predictive results and what the errors mean')
    figure(d,Path('weeks/week04/outputs/01_model_comparison.png'),'Figure 5.3  Model comparison carried forward from Week 4. Logistic regression was selected by development cross-validation before inspecting test predictions.',width=5.2)
    p(d,'Using a stratified 712/179 split, five development folds and fixed candidate settings, logistic regression achieved mean development ROC AUC 0.866. The tree achieved 0.853 and the baseline 0.500. Test accuracy was 82.1%, 77.7% and 61.5%, respectively. The final classification threshold was fixed at 0.50; the test results were not used to retune it.')
    figure(d,Path('weeks/week05/outputs/01_error_summary.png'),'Figure 5.4  The selected model makes 32 errors among 179 test passengers. Correct and incorrect predictions are shown as counts so the aggregate accuracy remains interpretable.',width=5.2)
    p(d,'The model detects 48 of 69 survivors, giving 69.6% recall, and 48 of 59 predicted survivors are correct, giving 81.4% precision. It detects all 45 female survivors but only 3 of 24 male survivors in this split. These differences limit what can be inferred from the overall accuracy. Earlier full-data exploration and possible family dependence also prevent this split from serving as independent external validation.')

    chapter(d,'6.  Strategic recommendations and success measures')
    p(d,'The strongest recommendations concern the quality of analysis and future evidence collection. Each action below is tied to an observed limitation and a measurable review criterion. Owners and timing are proposed responsibilities for a future project team; they are not commitments already agreed with an organization.')
    table(d,['Action and evidence','Proposed owner','Measure and timing'],[
        ['Maintain raw data, transformations and manifests; earlier duplicate and missingness choices changed interpretation.','Analyst','Every rerun passes dataset-hash and split checks; immediate.'],
        ['Include error counts and subgroup denominators; aggregate accuracy hides uneven recall.','Model reviewer','Review precision, recall and false-positive rates by relevant group; immediate.'],
        ['Acquire independent validation records; the full dataset informed prior exploration.','Research lead','Report external performance and uncertainty before broader use.'],
        ['Collect family IDs and relevant operational variables; current data omit dependence and mechanisms.','Data steward','Measure identifier coverage and linkage error before group-aware validation.'],
        ['Evaluate any proposed intervention separately; associations do not estimate intervention effects.','Domain specialist','Define outcomes, comparison design and decision criteria before a pilot.']],[2.7,1.1,2.3])
    sub(d,'Potential research and organizational value')
    p(d,'Versioned data and reproducible pipelines can reduce avoidable discrepancies during review and handover. Explicit subgroup evaluation can reveal weaknesses hidden by aggregate metrics. Better identifier and operational data can support more credible comparisons. These are expected process benefits; this project did not measure time savings, financial returns or improved safety outcomes.')
    p(d,'For a future transport or safety study, the historical class gradient could motivate investigation of access to assistance and route constraints. That would require contemporary operational data and specialist study design. The present model should not be used to rank people for rescue, deny services, or infer that one demographic group deserves priority.')

    chapter(d,'7.  Limitations and future improvement areas')
    p(d,'The data describe selected passenger records from a single historical event. They do not include every person aboard, and the sample is not a random draw from a modern population. Missingness, unrecorded context and shared family circumstances may affect both statistical and predictive results. Generalization beyond this setting is untested.')
    p(d,'The Week 3 tests rely on working independence and model assumptions; small p-values do not remove those assumptions. Simple age imputation is useful for a model pipeline but does not recover the uncertainty in missing values. The small age effect and the sensitivity result demonstrate why practical importance and robustness must accompany significance testing.')
    p(d,'The Week 4 comparison covers only fixed versions of logistic regression and a constrained decision tree. It does not prove that these are the best possible models. Cross-validation chooses between these candidates, while the internal test split provides a limited check. Prior exploration of the full dataset, related passengers across splits and small subgroup samples can all influence the apparent performance.')
    table(d,['Future work','Evidence needed before claiming improvement'],[
        ['Interactions and regularization','Pre-specified development comparison, followed by independent evaluation.'],
        ['Probability calibration','Out-of-fold calibration assessment and proper scoring rules on new data.'],
        ['Threshold optimization','Explicit error costs and development-only selection, with a new test evaluation.'],
        ['Family-aware validation','Reliable family or ticket linkage and group-separated partitions.'],
        ['Operational application','Relevant modern data, specialist oversight and evaluation of actual outcomes.']],[1.75,4.35])
    sub(d,'Workload plan and scope')
    p(d,'A proposed 32-hour Week 5 allocation is: evidence audit, 7 hours; synthesis and narrative, 7 hours; recommendations and limitations, 6 hours; figures and code documentation, 6 hours; and final review, 6 hours. This satisfies the planning range of 30–35 hours without representing an elapsed-time log.')

    chapter(d,'8.  Reproducibility and project handover')
    p(d,'The repository contains the source snapshot, shared data loader, weekly analysis scripts, Word reports, saved figures and machine-readable outputs. Weeks 1 and 2 retain the original supplied reports and clearly labelled reconstructed code. Weeks 3 and 4 provide executed analytical results; Week 5 reads these results without training another model or altering the test split.')
    code(d,'''# Run from the repository root after installing requirements.
python run_all.py
python -m unittest discover -s tests -v
# Rebuild Word reports separately, if required.
python build_remaining_reports.py''')
    p(d,'The Week 5 synthesis checks that Weeks 3 and 4 use the same dataset hash, then writes project_summary.json, recommendations.csv and the error-summary figure. The following excerpt shows how the selected model’s measured results are carried forward, avoiding manual re-entry or invented impact estimates.')
    code(d,'''statistical = json.loads(
    (ROOT/'weeks/week03/outputs/results.json').read_text())
modeling = json.loads(
    (ROOT/'weeks/week04/outputs/results.json').read_text())
assert statistical['sha256'] == modeling['dataset_sha256']
selected = next(r for r in modeling['metrics']
    if r['model'] == modeling['selected_model'])
tn, fp = modeling['confusion_matrix'][0]
fn, tp = modeling['confusion_matrix'][1]''')
    p(d,'The test suite checks passenger preservation, observed-age inference, split membership, saved predictions and reported confusion counts. Pinned dependencies and fixed random seeds support reproducibility. run_all.py regenerates analytical outputs; build_remaining_reports.py separately rebuilds the two Word reports using the approved Week 3 design, including the full-page pale blue cover.')
    sub(d,'Final conclusion')
    p(d,'The project demonstrates a complete learning workflow from raw data to evidence, predictions and recommendations. Its strongest outcome is a transparent and reviewable analysis that corrects earlier overstatements, explains model errors and separates observed associations from proposed actions. The next step is stronger evidence and validation, rather than treating a historical internal benchmark as a ready-to-use operational system.')
    references(d,[
        'Data source: https://github.com/mwaskom/seaborn-data/blob/master/titanic.csv',
        'Methods: https://scikit-learn.org/stable/common_pitfalls.html and https://scikit-learn.org/stable/modules/cross_validation.html',
        'Project sources: supplied Week 1 and Week 2 reports; Week 3 report and statistical outputs; Week 4 report, model metrics and test predictions.',
        'Repository: https://github.com/24f2001637/yuva-intern-Data-Science-with-Python-internship'])
    d.save(ROOT/'weeks/week05'/NAMES[5])

if __name__=='__main__':
    week4();week5()
    print('Built both Word reports.')
